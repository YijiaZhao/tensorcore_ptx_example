#!/usr/bin/env python3
"""Generate Tensor Core MMA summary PPT. Template: 40x22.5 inch slides, default 48pt font."""
from pptx import Presentation
from pptx.util import Pt
from lxml import etree
import os

prs = Presentation("nv_ppt_temp.pptx")

# Delete existing slides
for _ in range(len(prs.slides)):
    sldId = prs.slides._sldIdLst[0]
    rId = sldId.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    if rId:
        prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(sldId)

def slide9(section, title, items, sz=36):
    """Add Layout 9 slide: idx0=section header, idx13=body, idx14=subtitle(title)."""
    s = prs.slides.add_slide(prs.slide_layouts[9])
    for ph in s.placeholders:
        i = ph.placeholder_format.idx
        if i == 0: ph.text = section
        elif i == 14: ph.text = title
        elif i == 13:
            tf = ph.text_frame
            tf.clear()
            for j, item in enumerate(items):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = item
                p.font.size = Pt(sz)
    return s

# ===================== Slide 0: Cover =====================
s = prs.slides.add_slide(prs.slide_layouts[0])
for ph in s.placeholders:
    i = ph.placeholder_format.idx
    if i == 10: ph.text = "Blackwell Tensor Core\nMMA Instruction Deep Dive"
    elif i == 11: ph.text = "NVFP4 / MXFP4 on sm_100 & sm_120"
    elif i == 12: ph.text = "April 2026"

# ===================== Slide 1: Agenda =====================
slide9("AGENDA", "Agenda", [
    "1.  FP4 Data Format: e2m1 + Block Scaling",
    "2.  mma.sync on sm_120 (RTX 6000D)",
    "3.  tcgen05.mma on sm_100 (B100/B200)",
    "4.  sm_100 vs sm_120 Comparison",
    "5.  Verification Results",
    "6.  Key Takeaways",
], 40)

# ===================== Slide 2: FP4 Format =====================
slide9("DATA FORMAT", "FP4 e2m1 + Block Scaling", [
    "FP4 e2m1: 4-bit [sign | exp(2) | mant(1)], bias=1",
    "Values: 0, 0.5, 1.0, 1.5, 2.0, 3.0, 4.0, 6.0",
    "Packing: 2 FP4/byte, 8 FP4/uint32 (e.g. 0x22222222 = 8x1.0)",
    "",
    "Block Scaling: actual_value = scale x fp4_value",
    "  NVFP4 block16: ue4m3 scale per 16 elements (non-power-of-2)",
    "  MXFP4 block32: ue8m0 scale per 32 elements (power-of-2 only)",
], 36)

# ===================== Slide 3: Two Families =====================
slide9("OVERVIEW", "Two MMA Instruction Families", [
    "mma.sync  (sm_120, RTX 6000D)",
    "  32 threads (1 warp), tile M=16 N=8 K=64",
    "  Accumulator + A/B + Scale all in registers",
    "  8,192 FMA per instruction",
    "",
    "tcgen05.mma  (sm_100, B100/B200)",
    "  128 threads (4 warps, cta_group::1)",
    "  tile M=128 N=8 K=64, 1 thread issues (elect_one)",
    "  Accumulator + Scale in TMEM, A/B via smem descriptor",
    "  65,536 FMA per instruction (8x more)",
], 34)

# ===================== Slide 4: mma.sync PTX =====================
slide9("SM_120", "mma.sync NVFP4 PTX", [
    "Block32:  mma.sync.aligned.kind::mxf4nvf4",
    "   .block_scale.scale_vec::2X.m16n8k64.f32.e2m1.e2m1.f32.ue8m0",
    "",
    "Block16:  mma.sync.aligned.kind::mxf4nvf4",
    "   .block_scale.scale_vec::4X.m16n8k64.f32.e2m1.e2m1.f32.ue4m3",
    "",
    "Operands:  {d0-d3}, {a0-a3}, {b0-b1}, {c0-c3},",
    "           {sf_A}, {bid_a,tid_a}, {sf_B}, {bid_b,tid_b}",
    "All in registers. Every thread in warp executes.",
], 34)

# ===================== Slide 5: tcgen05.mma PTX =====================
slide9("SM_100", "tcgen05.mma NVFP4 PTX", [
    "NVFP4:  tcgen05.mma.cta_group::1",
    "   .kind::mxf4nvf4.block_scale.scale_vec::4X",
    "",
    "MXFP4:  tcgen05.mma.cta_group::1",
    "   .kind::mxf4.block_scale.scale_vec::2X",
    "",
    "Operands:  [tmem_c], desc_a, desc_b, idesc,",
    "           [tmem_sfa], [tmem_sfb], p",
    "TMEM + smem descriptors. Only 1 thread issues.",
], 34)

# ===================== Slide 6: tcgen05 Lifecycle =====================
slide9("SM_100", "tcgen05.mma Lifecycle", [
    "1. tcgen05.alloc    Allocate TMEM (whole warp)",
    "2. Fill smem         A/B FP4 data + scale factors",
    "3. tcgen05.st       Write scale to TMEM",
    "4. make_desc()      Construct SmemDescriptor (uint64)",
    "5. tcgen05.mma      Execute MMA (1 thread, elect_one)",
    "6. tcgen05.ld       Read result from TMEM to registers",
    "7. tcgen05.dealloc  Free TMEM (whole warp)",
], 36)

# ===================== Slide 7: Comparison =====================
slide9("COMPARISON", "sm_100 vs sm_120", [
    "                    sm_100 (B100)        sm_120 (6000D)",
    "Instruction     tcgen05.mma           mma.sync",
    "Threads          128 (4 warp)           32 (1 warp)",
    "Tile M            128                     16",
    "Accumulator     TMEM                   Registers",
    "A/B input       smem descriptor        Registers",
    "Launch          1 thread (elect)       32 (full warp)",
    "NVFP4 b16       mxf4nvf4 + 4X          mxf4nvf4 + 4X",
    "NVFP4 b32       N/A                    mxf4nvf4 + 2X",
    "MXFP4 b32       mxf4 + 2X              mxf4nvf4 + 2X",
], 32)

# ===================== Slide 8: Results =====================
slide9("RESULTS", "Verification Results", [
    "Test: A=B=FP4(1.0), scale=1.0 -> D = 64 x 1.0 = 64.0",
    "",
    "sm_120 mma.sync:",
    "  PASS  NVFP4 block32 + block16: 128/128 = 64.0",
    "  PASS  All 8 FP4 values + scale sweep + mixed scale",
    "",
    "sm_100 tcgen05.mma:",
    "  PASS  NVFP4 block16 (mxf4nvf4): 128/128 = 64.0",
    "  PASS  MXFP4 block32 (mxf4): 128/128 = 64.0",
    "  PASS  TMEM lifecycle + f16 baseline",
], 34)

# ===================== Slide 9: Takeaways =====================
slide9("TAKEAWAYS", "Key Takeaways", [
    "1. FP4 requires block scaling for correct results",
    "   Without scale: raw integer-like multiply (verified)",
    "",
    "2. sm_100 vs sm_120 architecturally different",
    "   sm_100: TMEM + smem descriptor + elect_one",
    "   sm_120: registers only (simpler, smaller tile)",
    "",
    "3. NVFP4 (mxf4nvf4) vs MXFP4 (mxf4): different PTX kind",
    "",
    "4. tcgen05.mma: alloc -> st -> mma -> ld -> dealloc",
    "",
    "5. Per-instruction FMA: sm_120=8K, sm_100=65K (8x)",
], 34)

out = os.path.join(os.path.dirname(os.path.abspath("nv_ppt_temp.pptx")), "tensor_core_mma_summary.pptx")
prs.save(out)
print(f"Saved: {out} ({len(prs.slides)} slides)")
